import random
import tkinter as tk
from dataclasses import dataclass, field
from datetime import datetime
from tkinter import filedialog, messagebox, ttk


# ============================================================================
# 1. DATOS BASE DEL MODELO
# ============================================================================
# Todos los tiempos se manejan internamente en minutos.
# La animación puede ser rápida, pero los relojes muestran el tiempo real.

MAX_DIAS = 6
BLOQUES_POR_BOLSA = 50
PASO_SIMULACION = 0.25

# Mediciones hechas por 1/4 de bolsa.
TIEMPO_MEZCLA_1_4 = 2 + 48.48 / 60
TIEMPO_TRANSPORTE_1_4 = 1 + 18.76 / 60

# Tiempos iniciales que puede modificar el usuario desde la interfaz.
TIEMPOS_INICIALES = {
    "mezcla": TIEMPO_MEZCLA_1_4 * 4,          # 11 min 13.92 s
    "transporte": TIEMPO_TRANSPORTE_1_4 * 4,  # 5 min 15.04 s
    "moldeo_dj": 36 + 14.87 / 60,             # 36 min 14.87 s
    "moldeo_vt": (36 + 14.87 / 60) * 2,       # Estimación para Víctor
    "hidratacion": 1.0,                       # 1 minuto
    "secado": 15 * 60,                        # 15 horas
    "retiro": 60.0,                           # 1 hora por lote
}

# Estados posibles de una bolsa dentro del Proceso 1.
ESPERA, MEZCLA, TRANSPORTE, MOLDEO, HIDRATACION = range(5)

# Posiciones horizontales de las estaciones en el Canvas (Proceso 1).
# Nota: hay una posición MÁS que etapas, porque "Terminada" (índice 5)
# es el destino visual final, no una etapa de proceso.
POSICIONES_X = [145, 380, 620, 900, 1130, 1280]


# ============================================================================
# 2. FUNCIONES AUXILIARES
# ============================================================================

def formatear_minutos(minutos: float, con_segundos: bool = True) -> str:
    """
    Convierte minutos decimales al formato HH:MM:SS.

    Ejemplo:
    36.25 minutos -> 00:36:15
    """
    segundos = max(0, round(minutos * 60))
    horas, segundos = divmod(segundos, 3600)
    mins, segundos = divmod(segundos, 60)

    if con_segundos:
        return f"{horas:02d}:{mins:02d}:{segundos:02d}"

    return f"{horas:02d}:{mins:02d}"


# ============================================================================
# 3. ESTRUCTURAS DE DATOS
# ============================================================================

@dataclass
class Bolsa:
    """
    Representa una bolsa dentro de la simulación.

    Una bolsa pasa por:
    Espera -> Mezcla -> Transporte -> Moldeo -> Hidratación -> Terminada.
    """

    numero: int
    cuadrilla: str
    etiqueta: str

    etapa: int = ESPERA
    tiempo_etapa: float = 0.0
    progreso: float = 0.0
    espera: float = 0.0
    terminada: bool = False  # NUEVO: marca que ya llegó a "Terminada"

    # Guarda los objetos gráficos de la bolsa dentro del Canvas.
    canvas_ids: list[int] = field(default_factory=list)


@dataclass
class EstadisticasCuadrilla:
    """
    Guarda resultados acumulados de cada cuadrilla.

    tiempo_laboral:
        Suma únicamente mezcla, transporte, moldeo e hidratación.

    tiempo_espera:
        Guarda el tiempo esperando la revolvedora.
    """

    bolsas: int = 0
    bloques: int = 0
    dias_trabajados: int = 0
    tiempo_laboral: float = 0.0
    tiempo_espera: float = 0.0

    @property
    def utilizacion(self) -> float:
        """
        Calcula el porcentaje del tiempo dedicado a trabajar
        respecto al tiempo laboral + tiempo de espera.
        """
        total = self.tiempo_laboral + self.tiempo_espera

        if total == 0:
            return 0

        return self.tiempo_laboral / total * 100


# ============================================================================
# 4. CLASE PRINCIPAL DE LA SIMULACIÓN
# ============================================================================

class SimulacionActual:
    """Interfaz gráfica y motor de la simulación."""

    COLORES = {
        "fondo": "#F0F4F8",
        "panel": "#FFFFFF",
        "azul": "#1A237E",
        "gris": "#546E7A",
        "don_jose": "#5D4037",
        "victor": "#6A1B9A",
        "naranja": "#EF6C00",
        "verde": "#2E7D32",
        "celeste": "#1565C0",
        "amarillo": "#F9A825",
        "rojo": "#C62828",
    }

    def __init__(self, root: tk.Tk):
        """Inicializa datos, ventana e interfaz."""

        self.root = root
        self.root.title(
            "D4 - Simulación actual | Industria Bloquera del Sureste"
        )
        self.root.geometry("1450x900")
        self.root.resizable(False, False)
        self.root.configure(bg=self.COLORES["fondo"])

        # Copia editable de los tiempos originales.
        self.config_tiempos = TIEMPOS_INICIALES.copy()

        # Control de animación.
        self._after_id = None
        self.velocidad = 1.0
        self.corriendo = False
        self.modo = "PRODUCCION"

        # Construye controles y deja la simulación lista.
        self.crear_interfaz()
        self.reiniciar(silencioso=True)

    # ========================================================================
    # 5. CREACIÓN DE LA INTERFAZ
    # ========================================================================

    def crear_interfaz(self):
        """Crea encabezado, indicadores, botones y Canvas."""

        c = self.COLORES

        # Encabezado principal.
        encabezado = tk.Frame(self.root, bg=c["azul"], height=64)
        encabezado.pack(fill="x")
        encabezado.pack_propagate(False)

        tk.Label(
            encabezado,
            text="INDUSTRIA BLOQUERA DEL SURESTE",
            bg=c["azul"],
            fg="white",
            font=("Arial", 17, "bold"),
        ).pack(pady=(7, 0))

        tk.Label(
            encabezado,
            text="D4 · Situación actual · Dos cuadrillas · Una revolvedora",
            bg=c["azul"],
            fg="#C5CAE9",
            font=("Arial", 9),
        ).pack()

        # Tarjetas de indicadores.
        indicadores = tk.Frame(self.root, bg=c["fondo"])
        indicadores.pack(fill="x", padx=10, pady=6)

        self.lbl_estado = self.crear_kpi(
            indicadores, "ESTADO", "LISTO", c["verde"]
        )

        self.lbl_reloj_dj = self.crear_kpi(
            indicadores,
            "DON JOSÉ · TIEMPO LABORAL",
            "00:00:00",
            c["don_jose"],
        )

        self.lbl_reloj_vt = self.crear_kpi(
            indicadores,
            "VÍCTOR · TIEMPO LABORAL",
            "00:00:00",
            c["victor"],
        )

        self.lbl_dia = self.crear_kpi(
            indicadores, "DÍA", "1 / 6", c["azul"]
        )

        self.lbl_rev = self.crear_kpi(
            indicadores, "REVOLVEDORA", "LIBRE", c["verde"]
        )

        self.lbl_almacen = self.crear_kpi(
            indicadores, "ALMACÉN", "0 bloques", c["verde"]
        )

        # Barra de botones.
        controles = tk.Frame(self.root, bg=c["fondo"])
        controles.pack(fill="x", padx=10, pady=2)

        ttk.Button(
            controles, text="▶ Iniciar", command=self.iniciar
        ).pack(side="left", padx=3)

        ttk.Button(
            controles, text="Ⅱ Pausar", command=self.pausar
        ).pack(side="left", padx=3)

        ttk.Button(
            controles, text="↻ Reiniciar", command=self.reiniciar
        ).pack(side="left", padx=3)

        ttk.Button(
            controles, text="➡ Proceso 2", command=self.iniciar_proceso2
        ).pack(side="left", padx=8)

        ttk.Button(
            controles, text="Resultados", command=self.mostrar_resultados
        ).pack(side="left", padx=3)

        ttk.Button(
            controles, text="Guardar Excel", command=self.guardar_excel
        ).pack(side="left", padx=3)

        ttk.Button(
            controles, text="⚙ Tiempos", command=self.ventana_tiempos
        ).pack(side="left", padx=3)

        tk.Label(
            controles,
            text="Velocidad:",
            bg=c["fondo"],
            font=("Arial", 9, "bold"),
        ).pack(side="left", padx=(18, 3))

        self.combo_velocidad = ttk.Combobox(
            controles,
            values=["Lenta", "Normal", "Rápida", "Muy rápida"],
            state="readonly",
            width=11,
        )
        self.combo_velocidad.set("Normal")
        self.combo_velocidad.bind(
            "<<ComboboxSelected>>",
            self.cambiar_velocidad,
        )
        self.combo_velocidad.pack(side="left")

        # Área principal donde se dibuja la simulación.
        self.canvas = tk.Canvas(
            self.root,
            width=1420,
            height=680,
            bg=c["panel"],
            highlightthickness=1,
            highlightbackground="#B0BEC5",
        )
        self.canvas.pack(padx=10, pady=5)

        # Aclaración sobre el significado de los relojes.
        tk.Label(
            self.root,
            text=(
                "Los relojes suman únicamente mezcla, transporte, "
                "moldeo e hidratación de cada cuadrilla."
            ),
            bg="#ECEFF1",
            fg=c["azul"],
            font=("Arial", 8),
        ).pack(fill="x", padx=10, pady=(0, 5))

    def crear_kpi(self, padre, titulo, valor, color):
        """Crea una tarjeta de indicador superior."""

        marco = tk.Frame(
            padre,
            bg="white",
            height=55,
            highlightbackground="#CFD8DC",
            highlightthickness=1,
        )
        marco.pack(side="left", padx=3, fill="x", expand=True)
        marco.pack_propagate(False)

        tk.Label(
            marco,
            text=titulo,
            bg="white",
            fg=self.COLORES["gris"],
            font=("Arial", 7, "bold"),
        ).pack(pady=(3, 0))

        etiqueta = tk.Label(
            marco,
            text=valor,
            bg="white",
            fg=color,
            font=("Arial", 10, "bold"),
        )
        etiqueta.pack()

        return etiqueta

    # ========================================================================
    # 6. PROCESO 1: DIBUJO DE PRODUCCIÓN
    # ========================================================================

    def dibujar_produccion(self):
        """Dibuja las dos líneas de producción y la revolvedora."""

        self.canvas.delete("all")
        c = self.COLORES

        self.canvas.create_text(
            710,
            18,
            text="PROCESO 1 · PRODUCCIÓN ACTUAL",
            fill=c["azul"],
            font=("Arial", 14, "bold"),
        )

        self.canvas.create_text(
            710,
            43,
            text=(
                "La animación es acelerada; los valores arriba "
                "de cada estación son tiempos reales."
            ),
            fill=c["gris"],
            font=("Arial", 10),
        )

        # Línea visual de Don José.
        self.dibujar_linea(
            y=175,
            titulo="CUADRILLA DON JOSÉ",
            color=c["don_jose"],
            fondo="#FFF8E1",
            objetivo=self.objetivo_dj,
        )

        # Línea visual de Víctor.
        self.dibujar_linea(
            y=430,
            titulo="CUADRILLA VÍCTOR",
            color=c["victor"],
            fondo="#F3E5F5",
            objetivo=self.objetivo_vt,
        )

        # Recurso compartido por ambas cuadrillas.
        self.canvas.create_rectangle(
            450,
            280,
            660,
            360,
            fill="#FFF3E0",
            outline=c["naranja"],
            width=3,
        )

        self.canvas.create_text(
            555,
            320,
            text="REVOLVEDORA\nCOMPARTIDA\n1 CUADRILLA A LA VEZ",
            fill=c["naranja"],
            font=("Arial", 9, "bold"),
            justify="center",
        )

        self.canvas.create_text(
            710,
            640,
            text=(
                "Tras la B3 se muestra la pausa de almuerzo. "
                "B3.5 no es una bolsa real ni afecta los resultados."
            ),
            fill=c["gris"],
            font=("Arial", 8),
        )

    def dibujar_linea(self, y, titulo, color, fondo, objetivo):
        """Dibuja estaciones, etiquetas y tiempos de una cuadrilla."""

        self.canvas.create_rectangle(
            25,
            y - 112,
            1395,
            y + 112,
            outline=color,
            width=2,
            fill=fondo,
        )

        self.canvas.create_text(
            40,
            y - 94,
            text=(
                f"{titulo} · Día {self.dia_actual} / {MAX_DIAS} "
                f"· Objetivo: {objetivo} bolsas"
            ),
            fill=color,
            font=("Arial", 11, "bold"),
            anchor="w",
        )

        # Cada tupla contiene:
        # posición horizontal, nombre del proceso y tiempo mostrado.
        estaciones = [
            (145, "Espera\nrevolvedora", "No suma reloj"),
            (
                380,
                "Mezcla",
                formatear_minutos(self.config_tiempos["mezcla"]),
            ),
            (
                620,
                "Carretilla",
                formatear_minutos(self.config_tiempos["transporte"]),
            ),
            (
                900,
                "Moldeo",
                formatear_minutos(
                    self.config_tiempos["moldeo_dj"]
                    if "JOSÉ" in titulo
                    else self.config_tiempos["moldeo_vt"]
                ),
            ),
            (
                1130,
                "Hidratación",
                formatear_minutos(self.config_tiempos["hidratacion"]),
            ),
            (1280, "Terminada", "Bolsa completa"),
        ]

        self.canvas.create_line(
            145,
            y,
            1280,
            y,
            fill=color,
            width=5,
        )

        # Dibuja cada estación con su nombre y tiempo.
        for x, nombre, tiempo in estaciones:
            self.canvas.create_oval(
                x - 8,
                y - 8,
                x + 8,
                y + 8,
                fill=color,
                outline="",
            )

            self.canvas.create_text(
                x,
                y - 48,
                text=tiempo,
                fill=color,
                font=("Arial", 8, "bold"),
            )

            self.canvas.create_text(
                x,
                y + 31,
                text=nombre,
                fill=self.COLORES["azul"],
                font=("Arial", 8),
                justify="center",
            )

    # ========================================================================
    # 7. CREACIÓN DE BOLSAS
    # ========================================================================

    def crear_bolsa(self, cuadrilla):
        """Crea una bolsa nueva si la cuadrilla puede continuar."""

        bolsa_actual = (
            self.bolsa_dj
            if cuadrilla == "Don José"
            else self.bolsa_vt
        )

        objetivo = (
            self.objetivo_dj
            if cuadrilla == "Don José"
            else self.objetivo_vt
        )

        hechas = (
            self.bolsas_hoy_dj
            if cuadrilla == "Don José"
            else self.bolsas_hoy_vt
        )

        # No crear otra bolsa si ya existe una activa,
        # si hay almuerzo o si ya se alcanzó el objetivo diario.
        if (
            bolsa_actual is not None
            or self.almuerzo_visible
            or hechas >= objetivo
            or objetivo == 0
        ):
            return

        numero = hechas + 1

        etiqueta = (
            f"B{numero}"
            if cuadrilla == "Don José"
            else f"B'{numero}"
        )

        bolsa = Bolsa(numero, cuadrilla, etiqueta)

        y = 175 if cuadrilla == "Don José" else 430

        color = (
            self.COLORES["don_jose"]
            if cuadrilla == "Don José"
            else self.COLORES["victor"]
        )

        # Rectángulo y texto que representan visualmente la bolsa.
        bolsa.canvas_ids = [
            self.canvas.create_rectangle(
                125,
                y - 12,
                165,
                y + 12,
                fill=color,
                outline="",
            ),
            self.canvas.create_text(
                145,
                y,
                text=etiqueta,
                fill="white",
                font=("Arial", 8, "bold"),
            ),
        ]

        if cuadrilla == "Don José":
            self.bolsa_dj = bolsa
        else:
            self.bolsa_vt = bolsa

    def estadistica(self, bolsa):
        """Devuelve el acumulador estadístico de la cuadrilla."""

        if bolsa.cuadrilla == "Don José":
            return self.est_dj

        return self.est_vt

    # ========================================================================
    # 8. MOTOR DE UNA BOLSA
    # ========================================================================

    def procesar_bolsa(self, bolsa, dt):
        """
        Avanza una bolsa según su etapa actual.

        dt representa el tiempo simulado transcurrido.
        Devuelve True si la bolsa terminó en este paso (llegó a Terminada).
        """

        # --------------------------------------------------------------------
        # Etapa 0: espera de revolvedora.
        # Este tiempo se guarda como espera y no suma al reloj laboral.
        # --------------------------------------------------------------------
        if bolsa.etapa == ESPERA:
            if not self.revolvedora_ocupada:
                self.revolvedora_ocupada = True
                self.quien_usa_revolvedora = bolsa.cuadrilla

                bolsa.etapa = MEZCLA
                bolsa.tiempo_etapa = 0.0
                bolsa.progreso = 0.0
            else:
                bolsa.espera += dt
                self.estadistica(bolsa).tiempo_espera += dt

            return False

        # Tiempo real de la etapa que está ejecutando la bolsa.
        tiempos = {
            MEZCLA: self.config_tiempos["mezcla"],
            TRANSPORTE: self.config_tiempos["transporte"],
            MOLDEO: (
                self.config_tiempos["moldeo_dj"]
                if bolsa.cuadrilla == "Don José"
                else self.config_tiempos["moldeo_vt"]
            ),
            HIDRATACION: self.config_tiempos["hidratacion"],
        }

        # Acumula tiempo de la etapa actual.
        bolsa.tiempo_etapa += dt

        # Calcula porcentaje de avance para mover la bolsa.
        bolsa.progreso = min(
            100,
            bolsa.tiempo_etapa / tiempos[bolsa.etapa] * 100,
        )

        # Mezcla, transporte, moldeo e hidratación sí suman al reloj laboral.
        self.estadistica(bolsa).tiempo_laboral += dt

        # También se registra cuánto se usa la revolvedora.
        if bolsa.etapa == MEZCLA:
            self.tiempo_revolvedora_ocupada += dt

        # Si la etapa sigue en curso, no se cambia todavía.
        if bolsa.tiempo_etapa < tiempos[bolsa.etapa]:
            return False

        # Al terminar mezcla, se libera la revolvedora.
        if bolsa.etapa == MEZCLA:
            self.revolvedora_ocupada = False
            self.quien_usa_revolvedora = None

        # Al terminar hidratación, se contabiliza la bolsa y llega a "Terminada".
        if bolsa.etapa == HIDRATACION:
            self.terminar_bolsa(bolsa)
            return True

        # Avanza a la siguiente etapa.
        bolsa.etapa += 1
        bolsa.tiempo_etapa = 0.0
        bolsa.progreso = 0.0
        return False

    def actualizar_visual_bolsa(self, bolsa):
        """Mueve gráficamente una bolsa según su avance."""

        if not bolsa.canvas_ids:
            return

        y = 175 if bolsa.cuadrilla == "Don José" else 430
        xs = POSICIONES_X

        # CORREGIDO: si ya llegó a "Terminada", se dibuja fija en esa
        # posición en vez de quedarse pegada en Hidratación.
        if bolsa.terminada:
            x = xs[5]
        else:
            x = xs[bolsa.etapa]

            # CORREGIDO: antes decía "< HIDRATACION", lo que excluía
            # justo la etapa de Hidratación y la bolsa nunca se
            # deslizaba hacia la estación "Terminada" (índice 5).
            if bolsa.etapa < len(xs) - 1:
                x += (
                    (xs[bolsa.etapa + 1] - x)
                    * bolsa.progreso
                    / 100
                )

        colores = {
            ESPERA: (
                self.COLORES["don_jose"]
                if bolsa.cuadrilla == "Don José"
                else self.COLORES["victor"]
            ),
            MEZCLA: self.COLORES["naranja"],
            TRANSPORTE: self.COLORES["celeste"],
            MOLDEO: self.COLORES["verde"],
            HIDRATACION: self.COLORES["amarillo"],
        }

        color_actual = (
            self.COLORES["verde"]
            if bolsa.terminada
            else colores[bolsa.etapa]
        )

        self.canvas.coords(
            bolsa.canvas_ids[0],
            x - 20,
            y - 12,
            x + 20,
            y + 12,
        )

        self.canvas.itemconfig(
            bolsa.canvas_ids[0],
            fill=color_actual,
        )

        self.canvas.coords(bolsa.canvas_ids[1], x, y)

    def terminar_bolsa(self, bolsa):
        """Cuenta una bolsa terminada y la agrega a la cola de secado."""

        bolsa.progreso = 100
        bolsa.terminada = True

        # El Proceso 2 usará esta lista para dibujar las eras.
        self.bolsas_terminadas.append(bolsa)

        est = self.estadistica(bolsa)

        est.bolsas += 1
        est.bloques += BLOQUES_POR_BOLSA

        # CORREGIDO: se programa el borrado del gráfico un poco después
        # de que se vea llegar a "Terminada", para que no se acumulen
        # rectángulos superpuestos en esa estación.
        ids_a_borrar = list(bolsa.canvas_ids)
        self.root.after(400, lambda ids=ids_a_borrar: self._borrar_graficos(ids))

        if bolsa.cuadrilla == "Don José":
            self.bolsas_hoy_dj += 1
            self.bolsa_dj = None

            # La tercera bolsa muestra pausa de almuerzo.
            if bolsa.numero == 3 and not self.almuerzo_mostrado:
                self.mostrar_almuerzo()
            else:
                self.crear_bolsa("Don José")

        else:
            self.bolsas_hoy_vt += 1
            self.bolsa_vt = None
            self.crear_bolsa("Víctor")

    def _borrar_graficos(self, ids):
        """Elimina del Canvas los gráficos de una bolsa ya contabilizada."""

        for id_grafico in ids:
            self.canvas.delete(id_grafico)

    # ========================================================================
    # 9. PAUSA DE ALMUERZO
    # ========================================================================

    def mostrar_almuerzo(self):
        """Muestra un aviso visual sin crear una B3.5 contable."""

        self.almuerzo_mostrado = True
        self.almuerzo_visible = True

        self.cartel = self.canvas.create_rectangle(
            480,
            270,
            940,
            355,
            fill="#FFECB3",
            outline=self.COLORES["naranja"],
            width=4,
        )

        self.texto_cartel = self.canvas.create_text(
            710,
            313,
            text="PAUSA DE ALMUERZO\nAmbas cuadrillas descansan",
            fill=self.COLORES["naranja"],
            font=("Arial", 14, "bold"),
            justify="center",
        )

        # Solo dura tres segundos visuales.
        self.root.after(3000, self.quitar_almuerzo)

    def quitar_almuerzo(self):
        """Quita el aviso y permite continuar con las bolsas."""

        self.canvas.delete(self.cartel)
        self.canvas.delete(self.texto_cartel)

        self.almuerzo_visible = False

        self.crear_bolsa("Don José")
        self.crear_bolsa("Víctor")

    # ========================================================================
    # 10. MOTOR PRINCIPAL DEL PROCESO 1
    # ========================================================================

    def iniciar(self):
        """Inicia la producción si todavía no termina."""

        if self.modo != "PRODUCCION":
            messagebox.showinfo(
                "Producción",
                "La producción ya terminó. Ejecute el Proceso 2.",
            )
            return

        if not self.corriendo:
            self.corriendo = True
            self.animar_produccion()

    def pausar(self):
        """Detiene temporalmente la simulación."""

        self.corriendo = False
        self.actualizar_indicadores()

    def animar_produccion(self):
        """Ejecuta un ciclo de animación y programa el siguiente."""

        if not self.corriendo or self.modo != "PRODUCCION":
            return

        dt = PASO_SIMULACION * self.velocidad

        # No se procesan bolsas mientras aparece el aviso de almuerzo.
        if not self.almuerzo_visible:
            self.crear_bolsa("Don José")
            self.crear_bolsa("Víctor")

            # Avanza una bolsa activa por cada cuadrilla.
            # CORREGIDO: antes solo se actualizaba el dibujo si la bolsa
            # seguía siendo la "activa" de la cuadrilla. Como al terminar
            # se pone self.bolsa_dj/self.bolsa_vt en None (o en una bolsa
            # nueva) DENTRO de procesar_bolsa, la bolsa que justo terminó
            # nunca recibía su última actualización visual y se quedaba
            # congelada en Hidratación. Ahora siempre se actualiza el
            # dibujo de la bolsa que se procesó en este ciclo.
            for bolsa in (self.bolsa_dj, self.bolsa_vt):
                if bolsa:
                    self.procesar_bolsa(bolsa, dt)
                    self.actualizar_visual_bolsa(bolsa)

            # Revisa si las dos cuadrillas concluyeron su objetivo diario.
            self.revisar_fin_dia()

        self.actualizar_indicadores()

        # Programa el siguiente ciclo visual.
        if self.corriendo:
            self._after_id = self.root.after(
                40,
                self.animar_produccion,
            )

    def revisar_fin_dia(self):
        """Comprueba si ambas cuadrillas terminaron el día actual."""

        dj_fin = (
            self.bolsas_hoy_dj >= self.objetivo_dj
            and self.bolsa_dj is None
        )

        vt_fin = (
            self.objetivo_vt == 0
            or (
                self.bolsas_hoy_vt >= self.objetivo_vt
                and self.bolsa_vt is None
            )
        )

        if not (dj_fin and vt_fin):
            return

        # Registra días trabajados.
        self.est_dj.dias_trabajados += 1

        if self.objetivo_vt:
            self.est_vt.dias_trabajados += 1

        # Si concluyó el día 6, se habilita Proceso 2.
        if self.dia_actual == MAX_DIAS:
            self.corriendo = False
            self.modo = "PRODUCCION_TERMINADA"
            self.actualizar_indicadores()

            messagebox.showinfo(
                "Producción terminada",
                (
                    "Se completaron los seis días. "
                    "Ya puede abrir el Proceso 2."
                ),
            )
            return

        # Si quedan días, prepara el siguiente.
        self.dia_actual += 1
        self.preparar_dia()
        self.dibujar_produccion()

    # ========================================================================
    # 11. PROCESO 2: SECADO, RETIRO Y ALMACÉN
    # ========================================================================

    def iniciar_proceso2(self):
        """Inicializa la simulación visual de secado y retiro."""

        if self.modo != "PRODUCCION_TERMINADA":
            messagebox.showinfo(
                "Proceso 2",
                "Primero complete los seis días de producción.",
            )
            return

        self.modo = "SECADO"

        # Reloj exclusivo del secado.
        self.tiempo_secado_actual = 0.0

        # Separa las bolsas terminadas por cuadrilla.
        self.cola_dj = [
            bolsa
            for bolsa in self.bolsas_terminadas
            if bolsa.cuadrilla == "Don José"
        ]

        self.cola_vt = [
            bolsa
            for bolsa in self.bolsas_terminadas
            if bolsa.cuadrilla == "Víctor"
        ]

        # Contadores de lotes retirados.
        self.retirado_dj = 0
        self.retirado_vt = 0

        # Estado del almacén.
        self.almacen_bloques = 0

        # Lote que el trabajador está retirando actualmente.
        self.lote_en_retiro = None
        self.tiempo_retiro_actual = 0.0

        # Sirve para alternar entre eras.
        self.turno_retiro = "Don José"

        self.dibujar_proceso2()
        self.animar_proceso2()

    def animar_proceso2(self):
        """Controla secado y retiro de bloques."""

        if self.modo != "SECADO":
            return

        dt = PASO_SIMULACION * self.velocidad

        # Fase A: secado.
        if self.tiempo_secado_actual < self.config_tiempos["secado"]:
            self.tiempo_secado_actual = min(
                self.config_tiempos["secado"],
                self.tiempo_secado_actual + dt,
            )

        # Fase B: retiro de bloques.
        else:
            self.avanzar_retiro(dt)

        self.dibujar_proceso2()
        self.actualizar_indicadores()

        if self.modo == "SECADO":
            self._after_id = self.root.after(
                40,
                self.animar_proceso2,
            )

    def siguiente_lote(self):
        """
        Elige el siguiente lote seco por retirar.

        Alterna entre Don José y Víctor cuando ambos tienen bolsas
        pendientes; así se ve que el personal atiende ambas eras.
        """

        preferida = (
            self.cola_dj
            if self.turno_retiro == "Don José"
            else self.cola_vt
        )

        alternativa = (
            self.cola_vt
            if self.turno_retiro == "Don José"
            else self.cola_dj
        )

        if preferida:
            bolsa = preferida.pop(0)

        elif alternativa:
            bolsa = alternativa.pop(0)

        else:
            return None

        self.turno_retiro = (
            "Víctor"
            if bolsa.cuadrilla == "Don José"
            else "Don José"
        )

        return bolsa

    def avanzar_retiro(self, dt):
        """Mueve un lote desde su era al almacén."""

        # Si no hay lote en movimiento, se toma el siguiente.
        if self.lote_en_retiro is None:
            self.lote_en_retiro = self.siguiente_lote()
            self.tiempo_retiro_actual = 0.0

            # Si no hay lotes, concluye el Proceso 2.
            if self.lote_en_retiro is None:
                self.modo = "PROCESO2_TERMINADO"

                messagebox.showinfo(
                    "Proceso 2 terminado",
                    f"Bloques almacenados: {self.almacen_bloques}",
                )
                return

        # Avanza el retiro del lote actual.
        self.tiempo_retiro_actual += dt

        # Si termina, suma 50 bloques al almacén.
        if self.tiempo_retiro_actual >= self.config_tiempos["retiro"]:
            if self.lote_en_retiro.cuadrilla == "Don José":
                self.retirado_dj += 1
            else:
                self.retirado_vt += 1

            self.almacen_bloques += BLOQUES_POR_BOLSA
            self.lote_en_retiro = None

    def dibujar_proceso2(self):
        """Dibuja eras, secado, retiro y almacén."""

        self.canvas.delete("all")
        c = self.COLORES

        self.canvas.create_text(
            710,
            20,
            text="PROCESO 2 · SECADO, RETIRO Y ALMACÉN",
            fill=c["azul"],
            font=("Arial", 14, "bold"),
        )

        # Texto superior de estado del secado.
        if self.tiempo_secado_actual < self.config_tiempos["secado"]:
            porcentaje = (
                self.tiempo_secado_actual
                / self.config_tiempos["secado"]
                * 100
            )

            estado = (
                f"SECANDO · {porcentaje:.1f}% · "
                f"{formatear_minutos(self.tiempo_secado_actual)} / "
                f"{formatear_minutos(self.config_tiempos['secado'])}"
            )

        else:
            estado = "SECADO TERMINADO · PERSONAL RETIRANDO BLOQUES"

        self.canvas.create_text(
            710,
            50,
            text=estado,
            fill=c["celeste"],
            font=("Arial", 10, "bold"),
        )

        # Total de bolsas producidas por cuadrilla.
        total_dj = self.est_dj.bolsas
        total_vt = self.est_vt.bolsas

        # Bolsas que todavía están en la era.
        pendientes_dj = len(self.cola_dj) + (
            1
            if (
                self.lote_en_retiro
                and self.lote_en_retiro.cuadrilla == "Don José"
            )
            else 0
        )

        pendientes_vt = len(self.cola_vt) + (
            1
            if (
                self.lote_en_retiro
                and self.lote_en_retiro.cuadrilla == "Víctor"
            )
            else 0
        )

        # Era de Don José.
        self.dibujar_era(
            105,
            115,
            "ERA DON JOSÉ",
            c["don_jose"],
            total_dj,
            pendientes_dj,
            self.retirado_dj,
        )

        # Era de Víctor.
        self.dibujar_era(
            705,
            115,
            "ERA VÍCTOR",
            c["victor"],
            total_vt,
            pendientes_vt,
            self.retirado_vt,
        )

        # Área visual del personal de retiro.
        self.canvas.create_rectangle(
            505,
            405,
            915,
            505,
            fill="#FFF3E0",
            outline=c["naranja"],
            width=3,
        )

        self.canvas.create_text(
            710,
            430,
            text="PERSONAL DE RETIRO",
            fill=c["naranja"],
            font=("Arial", 12, "bold"),
        )

        self.canvas.create_text(
            710,
            462,
            text="👷 Retira un lote de 50 bloques de cada era por turnos",
            fill=c["gris"],
            font=("Arial", 10),
        )

        # Si hay lote en movimiento, se anima entre era y almacén.
        if self.lote_en_retiro:
            origen = (
                360
                if self.lote_en_retiro.cuadrilla == "Don José"
                else 1060
            )

            progreso = min(
                1,
                self.tiempo_retiro_actual
                / self.config_tiempos["retiro"],
            )

            x = origen + (710 - origen) * progreso

            self.canvas.create_rectangle(
                x - 30,
                520,
                x + 30,
                555,
                fill="#8D6E63",
                outline="",
            )

            self.canvas.create_text(
                x,
                537,
                text=self.lote_en_retiro.etiqueta,
                fill="white",
                font=("Arial", 9, "bold"),
            )

            self.canvas.create_text(
                710,
                580,
                text=(
                    f"Retirando {self.lote_en_retiro.etiqueta}: "
                    f"{progreso * 100:.1f}% · "
                    f"{formatear_minutos(self.tiempo_retiro_actual)} / "
                    f"{formatear_minutos(self.config_tiempos['retiro'])}"
                ),
                fill=c["naranja"],
                font=("Arial", 9, "bold"),
            )

        else:
            self.canvas.create_text(
                710,
                580,
                text="Esperando lote seco para retirar",
                fill=c["gris"],
                font=("Arial", 9),
            )

        # Área del almacén.
        self.canvas.create_rectangle(
            1000,
            405,
            1350,
            605,
            fill="#E8F5E9",
            outline=c["verde"],
            width=3,
        )

        self.canvas.create_text(
            1175,
            435,
            text="ALMACÉN",
            fill=c["verde"],
            font=("Arial", 14, "bold"),
        )

        self.canvas.create_text(
            1175,
            470,
            text=f"{self.almacen_bloques} bloques",
            fill=c["azul"],
            font=("Arial", 16, "bold"),
        )

        # Dibuja bloques acumulados visualmente en el almacén.
        for i in range(min(24, self.almacen_bloques // BLOQUES_POR_BOLSA)):
            x = 1040 + (i % 8) * 33
            y = 515 + (i // 8) * 25

            self.canvas.create_rectangle(
                x,
                y,
                x + 25,
                y + 16,
                fill="#78909C",
                outline="#455A64",
            )

    def dibujar_era(self, x, y, titulo, color, total, pendientes, retirados):
        """Dibuja una era con bloques visibles."""

        self.canvas.create_rectangle(
            x,
            y,
            x + 500,
            y + 230,
            fill="#F8F9FA",
            outline=color,
            width=3,
        )

        self.canvas.create_text(
            x + 250,
            y + 25,
            text=titulo,
            fill=color,
            font=("Arial", 12, "bold"),
        )

        self.canvas.create_text(
            x + 250,
            y + 52,
            text=(
                f"Producidas: {total} bolsas · "
                f"En era: {pendientes} · "
                f"Retiradas: {retirados}"
            ),
            fill=self.COLORES["gris"],
            font=("Arial", 9),
        )

        # Representación compacta de bloques.
        for i in range(min(48, pendientes * 4)):
            bx = x + 35 + (i % 12) * 36
            by = y + 88 + (i // 12) * 27

            self.canvas.create_rectangle(
                bx,
                by,
                bx + 28,
                by + 18,
                fill="#B0BEC5",
                outline="#607D8B",
            )

        # Barra visual del progreso de secado.
        if self.tiempo_secado_actual < self.config_tiempos["secado"]:
            progreso = (
                self.tiempo_secado_actual
                / self.config_tiempos["secado"]
            )

            self.canvas.create_rectangle(
                x + 28,
                y + 205,
                x + 472,
                y + 217,
                fill="#E3F2FD",
                outline="",
            )

            self.canvas.create_rectangle(
                x + 28,
                y + 205,
                x + 28 + 444 * progreso,
                y + 217,
                fill=self.COLORES["celeste"],
                outline="",
            )

    # ========================================================================
    # 12. INDICADORES Y CONFIGURACIÓN
    # ========================================================================

    def actualizar_indicadores(self):
        """Actualiza las tarjetas superiores."""

        self.lbl_reloj_dj.config(
            text=formatear_minutos(self.est_dj.tiempo_laboral)
        )

        self.lbl_reloj_vt.config(
            text=formatear_minutos(self.est_vt.tiempo_laboral)
        )

        self.lbl_dia.config(
            text=f"{self.dia_actual} / {MAX_DIAS}"
        )

        self.lbl_almacen.config(
            text=f"{getattr(self, 'almacen_bloques', 0)} bloques"
        )

        # Estado del recurso compartido.
        if self.revolvedora_ocupada:
            self.lbl_rev.config(
                text=f"OCUPADA\n{self.quien_usa_revolvedora}",
                fg=self.COLORES["rojo"],
            )
        else:
            self.lbl_rev.config(
                text="LIBRE",
                fg=self.COLORES["verde"],
            )

        # Estado principal de toda la simulación.
        estado = "EN EJECUCIÓN" if self.corriendo else "PAUSADA / LISTO"
        color = (
            self.COLORES["verde"]
            if self.corriendo
            else self.COLORES["naranja"]
        )

        if self.modo == "PRODUCCION_TERMINADA":
            estado = "PRODUCCIÓN TERMINADA"
            color = self.COLORES["verde"]

        elif self.modo == "SECADO":
            estado = "SECADO / RETIRO"
            color = self.COLORES["celeste"]

        elif self.modo == "PROCESO2_TERMINADO":
            estado = "PROCESO 2 TERMINADO"
            color = self.COLORES["verde"]

        self.lbl_estado.config(text=estado, fg=color)

    def cambiar_velocidad(self, _evento=None):
        """Cambia la rapidez visual de la simulación."""

        valores = {
            "Lenta": 0.35,
            "Normal": 1,
            "Rápida": 3.5,
            "Muy rápida": 9,
        }

        self.velocidad = valores.get(
            self.combo_velocidad.get(),
            1,
        )

    def preparar_dia(self):
        """Crea los objetivos de producción del día."""

        self.objetivo_dj = random.choice([6, 7, 8])

        if self.plan_victor[self.dia_actual - 1]:
            self.objetivo_vt = random.choice([3, 4])
        else:
            self.objetivo_vt = 0

        self.bolsas_hoy_dj = 0
        self.bolsas_hoy_vt = 0

        self.bolsa_dj = None
        self.bolsa_vt = None

        self.almuerzo_mostrado = False
        self.almuerzo_visible = False

    def reiniciar(self, silencioso=False):
        """Restablece todos los datos y genera un nuevo escenario."""

        self.corriendo = False

        # Cancela una animación activa, si existe.
        if self._after_id:
            try:
                self.root.after_cancel(self._after_id)
            except tk.TclError:
                pass

        self._after_id = None

        self.modo = "PRODUCCION"
        self.dia_actual = 1

        # Víctor trabaja aleatoriamente 4 o 5 de los 6 días.
        self.plan_victor = [False] * MAX_DIAS

        for indice in random.sample(
            range(MAX_DIAS),
            random.choice([4, 5]),
        ):
            self.plan_victor[indice] = True

        # Reinicia estadísticas.
        self.est_dj = EstadisticasCuadrilla()
        self.est_vt = EstadisticasCuadrilla()

        self.bolsas_terminadas = []

        self.revolvedora_ocupada = False
        self.quien_usa_revolvedora = None
        self.tiempo_revolvedora_ocupada = 0.0

        self.almacen_bloques = 0

        self.preparar_dia()
        self.dibujar_produccion()
        self.actualizar_indicadores()

        if not silencioso:
            messagebox.showinfo(
                "Reinicio",
                "Se creó un nuevo escenario de seis días.",
            )

    # ========================================================================
    # 13. RESULTADOS, CONFIGURACIÓN DE TIEMPOS Y EXCEL
    # ========================================================================

    def mostrar_resultados(self):
        """Abre una ventana con los resultados principales."""

        total_bolsas = self.est_dj.bolsas + self.est_vt.bolsas
        total_bloques = self.est_dj.bloques + self.est_vt.bloques

        datos = (
            "RESULTADOS · ESCENARIO ACTUAL\n\n"
            f"Don José: {self.est_dj.bolsas} bolsas · "
            f"{self.est_dj.bloques} bloques\n"
            f"  Tiempo laboral acumulado: "
            f"{formatear_minutos(self.est_dj.tiempo_laboral)}\n"
            f"  Espera por revolvedora: "
            f"{formatear_minutos(self.est_dj.tiempo_espera)}\n"
            f"  Utilización: {self.est_dj.utilizacion:.1f}%\n\n"
            f"Víctor: {self.est_vt.bolsas} bolsas · "
            f"{self.est_vt.bloques} bloques\n"
            f"  Tiempo laboral acumulado: "
            f"{formatear_minutos(self.est_vt.tiempo_laboral)}\n"
            f"  Espera por revolvedora: "
            f"{formatear_minutos(self.est_vt.tiempo_espera)}\n"
            f"  Utilización: {self.est_vt.utilizacion:.1f}%\n\n"
            f"TOTAL: {total_bolsas} bolsas · {total_bloques} bloques\n"
            f"Almacén: {getattr(self, 'almacen_bloques', 0)} bloques\n"
            f"Revolvedora ocupada: "
            f"{formatear_minutos(self.tiempo_revolvedora_ocupada)}"
        )

        ventana = tk.Toplevel(self.root)
        ventana.title("Resultados - Escenario actual")
        ventana.geometry("650x520")

        tk.Label(
            ventana,
            text=datos,
            justify="left",
            anchor="w",
            padx=25,
            pady=25,
            font=("Consolas", 10),
            bg="white",
            fg=self.COLORES["azul"],
        ).pack(fill="both", expand=True)

    def ventana_tiempos(self):
        """Permite editar los tiempos sin modificar el código."""

        ventana = tk.Toplevel(self.root)
        ventana.title("Tiempos del modelo")
        ventana.geometry("550x450")

        marco = tk.Frame(ventana, padx=20, pady=20)
        marco.pack(fill="both", expand=True)

        campos = [
            ("mezcla", "Preparación de mezcla por bolsa"),
            ("transporte", "Carretilla por bolsa"),
            ("moldeo_dj", "Moldeo Don José"),
            ("moldeo_vt", "Moldeo Víctor"),
            ("hidratacion", "Hidratación"),
            ("secado", "Secado"),
            ("retiro", "Retiro por lote"),
        ]

        entradas = {}

        for fila, (clave, texto) in enumerate(campos):
            tk.Label(
                marco,
                text=texto,
            ).grid(row=fila, column=0, sticky="w", pady=6)

            entrada = tk.Entry(marco, width=16)
            entrada.insert(
                0,
                f"{self.config_tiempos[clave]:.4f}",
            )
            entrada.grid(row=fila, column=1, padx=10)

            tk.Label(
                marco,
                text="minutos",
            ).grid(row=fila, column=2, sticky="w")

            entradas[clave] = entrada

        def guardar():
            """Valida y guarda los nuevos tiempos escritos."""

            try:
                valores = {
                    clave: float(entrada.get())
                    for clave, entrada in entradas.items()
                }

                if any(valor <= 0 for valor in valores.values()):
                    raise ValueError

                self.config_tiempos.update(valores)
                ventana.destroy()

                # Redibuja los tiempos visibles de Proceso 1.
                if self.modo == "PRODUCCION":
                    self.dibujar_produccion()

                messagebox.showinfo(
                    "Tiempos",
                    "Los tiempos se actualizaron.",
                )

            except ValueError:
                messagebox.showerror(
                    "Dato inválido",
                    "Escriba valores numéricos mayores que cero.",
                )

        ttk.Button(
            marco,
            text="Guardar tiempos",
            command=guardar,
        ).grid(
            row=len(campos),
            column=0,
            columnspan=3,
            pady=20,
        )

    def guardar_excel(self):
        """Exporta resultados principales a un archivo de Excel (.xlsx)."""

        archivo = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Archivo de Excel", "*.xlsx")],
        )

        if not archivo:
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            messagebox.showerror(
                "Falta openpyxl",
                "Para guardar en Excel necesitas instalar la librería "
                "openpyxl.\n\nEjecuta en la terminal:\n"
                "pip install openpyxl",
            )
            return

        try:
            libro = Workbook()
            hoja = libro.active
            hoja.title = "Resultados"

            hoja.append(["Indicador", "Valor"])

            relleno_encabezado = PatternFill(
                start_color="1A237E",
                end_color="1A237E",
                fill_type="solid",
            )
            fuente_encabezado = Font(color="FFFFFF", bold=True)

            for celda in hoja[1]:
                celda.fill = relleno_encabezado
                celda.font = fuente_encabezado
                celda.alignment = Alignment(horizontal="center")

            filas = [
                ["Fecha", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
                ["Bolsas Don José", self.est_dj.bolsas],
                ["Bolsas Víctor", self.est_vt.bolsas],
                ["Bloques Don José", self.est_dj.bloques],
                ["Bloques Víctor", self.est_vt.bloques],
                [
                    "Tiempo laboral Don José",
                    formatear_minutos(self.est_dj.tiempo_laboral),
                ],
                [
                    "Tiempo laboral Víctor",
                    formatear_minutos(self.est_vt.tiempo_laboral),
                ],
                [
                    "Espera Don José (min)",
                    round(self.est_dj.tiempo_espera, 2),
                ],
                [
                    "Espera Víctor (min)",
                    round(self.est_vt.tiempo_espera, 2),
                ],
                [
                    "Utilización Don José (%)",
                    round(self.est_dj.utilizacion, 1),
                ],
                [
                    "Utilización Víctor (%)",
                    round(self.est_vt.utilizacion, 1),
                ],
                [
                    "Revolvedora ocupada (min)",
                    round(self.tiempo_revolvedora_ocupada, 2),
                ],
                ["Almacén (bloques)", getattr(self, "almacen_bloques", 0)],
            ]

            for fila in filas:
                hoja.append(fila)

            hoja.column_dimensions["A"].width = 30
            hoja.column_dimensions["B"].width = 25

            libro.save(archivo)

            messagebox.showinfo(
                "Excel",
                f"Resultados guardados en:\n{archivo}",
            )

        except OSError as error:
            messagebox.showerror(
                "Error",
                f"No se pudo guardar el archivo:\n{error}",
            )


# ============================================================================
# 14. PUNTO DE ENTRADA
# ============================================================================

if __name__ == "__main__":
    raiz = tk.Tk()
    app = SimulacionActual(raiz)
    raiz.mainloop()